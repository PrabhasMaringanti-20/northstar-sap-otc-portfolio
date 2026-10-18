"""Create offline viewer, editable workbook/report, charts and ten-slide presentation."""
from pathlib import Path
from html import escape
import csv,json,re,shutil,subprocess,os,tempfile
os.environ.setdefault('MPLCONFIGDIR',str(Path(tempfile.gettempdir())/'northstar-matplotlib'))
os.environ.setdefault('XDG_CACHE_HOME',str(Path(tempfile.gettempdir())/'northstar-cache'))
Path(os.environ['XDG_CACHE_HOME']).mkdir(parents=True,exist_ok=True)
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Inches,Pt
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'downloads';OUT.mkdir(exist_ok=True)
cat=json.loads((ROOT/'quality/catalog.json').read_text());res=json.loads((ROOT/'phase-08/results.json').read_text());ctx=res['context']
def csvread(p):
 with (ROOT/p).open(newline='',encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def graph(source,target):
 node=r'(\w+)(?:\[([^\]]+)\]|\{([^}]+)\})?'
 pattern=re.compile('^'+node+r'\s*(-->|-- (.*?) -->|-\. (.*?) \.->)\s*'+node+'$')
 lines=['digraph process {','graph [rankdir=TB,bgcolor="white",pad="0.25",nodesep="0.3",ranksep="0.45"];','node [shape=box,style="rounded,filled",fillcolor="#eaf3f7",color="#397187",fontname="DejaVu Sans",fontsize=11];','edge [color="#607887",fontname="DejaVu Sans",fontsize=10];']
 for raw in source.read_text().splitlines()[1:]:
  m=pattern.match(raw.strip())
  if not m: raise ValueError('Unparsed flowchart edge: '+raw)
  a,al,ad,arrow,lab1,lab2,b,bl,bd=m.groups()
  for ident,label,decision in [(a,al,ad),(b,bl,bd)]:
   if label or decision:lines.append(f'{ident} [label={json.dumps(label or decision)},shape={"diamond" if decision else "box"}];')
  lines.append(f'{a} -> {b} [label={json.dumps(lab1 or lab2 or "")},style={"dashed" if lab2 else "solid"}];')
 lines.append('}');dot=target.with_suffix('.dot');dot.write_text('\n'.join(lines))
 if shutil.which('dot'):subprocess.run(['dot','-Tsvg',str(dot),'-o',str(target)],check=True)
 elif not target.exists():raise RuntimeError('Graphviz unavailable and no supplied SVG')
 text=target.read_text();return text[text.index('<svg'):]
maps={'asis':graph(ROOT/'phase-03/as-is.mmd',ROOT/'phase-03/as-is.svg'),'tobe':graph(ROOT/'phase-05/to-be.mmd',ROOT/'phase-05/to-be.svg')}

# Exportable figure; separate axes avoid comparing days to percentages.
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'svg.fonttype':'none'})
fig,axs=plt.subplots(1,2,figsize=(12,4.3),gridspec_kw={'width_ratios':[1,1.5]})
days=[m for m in res['metrics'] if m['unit']=='calendar days'];pct=[m for m in res['metrics'] if m['unit']=='%']
for ax,ms,limit,label in [(axs[0],days,8,'Calendar days'),(axs[1],pct,100,'Percent')]:
 vals=[m['value'] for m in ms]; ax.barh([m['name'] for m in ms],vals,color='#087e8b',height=.5);ax.invert_yaxis();ax.set_xlim(0,limit);ax.set_xlabel(label);ax.spines[['top','right']].set_visible(False)
 for i,v in enumerate(vals):ax.text(v+.15,i,f'{v:.2f}',va='center',fontsize=9)
fig.suptitle('Synthetic baseline — 300 orders | snapshot 11 October 2026',fontsize=14,fontweight='bold')
fig.text(.02,.01,'Illustrative generated data, not measured business performance. Definitions and denominator exclusions are in the KPI contract.',fontsize=8)
fig.tight_layout(rect=[0,.05,1,.93]);fig.savefig(ROOT/'phase-08/kpi-baseline.svg');fig.savefig(ROOT/'phase-08/kpi-baseline.png',dpi=160);plt.close(fig)
chart=(ROOT/'phase-08/kpi-baseline.svg').read_text();chart=chart[chart.index('<svg'):]

def clean(t):return re.sub(r'\[([^]]+)\]\([^)]+\)',r'\1',t).replace('**','').replace('`','')
def cells(t):return [clean(c.strip()) for c in t.strip().strip('|').split('|')]
def sep(t):return all(re.fullmatch(r':?-+:?',x.strip()) for x in t.strip().strip('|').split('|'))
def inline(t):
 t=escape(t)
 t=re.sub(r'\*\*([^*]+)\*\*',r'<strong>\1</strong>',t)
 t=re.sub(r'`([^`]+)`',r'<code>\1</code>',t)
 # Library reads offline: textual file references remain text; source register has clickable external links.
 return re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1 <small>(\2)</small>',t)
def mdhtml(p):
 lines=p.read_text().splitlines();out=[];i=0
 while i<len(lines):
  line=lines[i]
  if line.startswith('```'):
   lang=line[3:];block=[];i+=1
   while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
   if lang=='mermaid' and p.parent.name in ['phase-03','phase-05']:out.append('<div class="diagram">'+maps['asis' if p.parent.name=='phase-03' else 'tobe']+'</div>')
   else:out.append('<pre>'+escape('\n'.join(block))+'</pre>')
  elif line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].startswith('|'):
    if not sep(lines[i]):rows.append(cells(lines[i]))
    i+=1
   out.append('<div class="scroll"><table><thead><tr>'+''.join('<th>'+escape(v)+'</th>' for v in rows[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+escape(v)+'</td>' for v in row)+'</tr>' for row in rows[1:])+'</tbody></table></div>');continue
  elif line.startswith('#'):
   level=min(len(line)-len(line.lstrip('#')),4);out.append(f'<h{level}>'+inline(line.lstrip('# '))+f'</h{level}>')
  elif line.startswith('>'):out.append('<blockquote>'+inline(line.lstrip('> '))+'</blockquote>')
  elif line.strip():out.append('<p>'+inline(line)+'</p>')
  i+=1
 return '\n'.join(out)

docpaths=[ROOT/p for p in ['README.md','PROJECT_GUIDE.md','RUN_LOCALLY.md','GITHUB_SETUP.md'] if (ROOT/p).exists()]+sorted(ROOT.glob('phase-*/*.md'))+sorted(ROOT.glob('registers/*.md'))+sorted(ROOT.glob('quality/*.md'))
docs=[dict(path=str(p.relative_to(ROOT)),html=mdhtml(p)) for p in docpaths]
orders=csvread('phase-08/data/order_summary.csv');customers={c['customer_id']:c['customer_name'] for c in csvread('phase-08/data/customers.csv')};products={p['product_id']:p['product_name'] for p in csvread('phase-08/data/products.csv')}
for o in orders:
 for k in ['ordered_qty','unit_price_pence','net_value_pence','cancelled','partial_allowed','shipped_qty','received_qty']:o[k]=int(o[k])
 o['customer_name']=customers[o['customer_id']];o['product_name']=products[o['product_id']]
data=dict(results=res,catalog=cat,orders=orders,docs=docs,maps=maps,chart=chart,kpis=csvread('phase-08/kpi-contract.csv'))
for name,numeric in [('deliveries',['shipped_qty']),('invoices',['net_amount_pence']),('allocations',['allocated_pence'])]:
 data[name]=csvread(f'phase-08/data/{name}.csv')
 for row in data[name]:
  for f in numeric:row[f]=int(row[f])
payload=json.dumps(data,ensure_ascii=False).replace('</',r'<\/')
nav=[('overview','Overview'),('orders','Orders'),('process','Process maps'),('requirements','Requirements'),('analysis','Data & KPIs'),('uat','UAT & checks'),('evidence','SAP evidence'),('library','Document library')]
html='''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>NorthStar | Order-to-Cash Portfolio</title><meta name="description" content="Independent simulated business analysis portfolio with synthetic data and unexecuted SAP UAT."><style>'''+(ROOT/'viewer/style.css').read_text()+'''</style></head><body><a class="skip" href="#content">Skip to content</a><aside><div class="brand"><span>✦</span> NorthStar</div><div class="subbrand">ORDER-TO-CASH · ANALYSIS PORTFOLIO</div><nav aria-label="Portfolio sections">'''+''.join(f'<button data-view="{v}">{label}</button>' for v,label in nav)+'''</nav><footer>Independent simulated project<br>Offline · No SAP connection<br>Free-tool documentation route</footer></aside><main><div class="topbar"><span>BUSINESS ANALYSIS / NORTHSTAR UK</span><span class="badge">SYNTHETIC DATA · 11 OCT 2026</span></div><div id="content" tabindex="-1"></div></main><script>const DATA='''+payload+';</script><script>'+(ROOT/'viewer/app.js').read_text()+'</script></body></html>'
(ROOT/'viewer/index.html').write_text(html,encoding='utf-8')
shutil.copyfile(ROOT/'viewer/index.html',ROOT/'index.html')
shutil.copyfile(ROOT/'viewer/index.html',OUT/'NorthStar_Offline_Portfolio.html')

# Editable master workbook, including complete catalogues and raw data.
wb=Workbook();ws=wb.active;ws.title='Read Me'
for row in [['NorthStar complete portfolio','Independent simulation · 11 October 2026'],['Evidence boundary','22 SAP UAT cases Not executed; synthetic data checks executed separately'],['Currency','GBP integer pence; all finance net-only; VAT omitted'],['Reproduction','python scripts/build_all.py; source CSV/Markdown is authoritative'],['Calculation behaviour','Workbook contains generated values, not live KPI formulas. Rebuild Python/SQL outputs after source edits; see KPI Contract.'],['Free editing','LibreOffice Calc; no paid Microsoft Office licence required'],['Learner status','Understanding and personal contribution review pending']]:ws.append(row)
sheets=[('Requirements','phase-04/requirements.csv'),('Rules','phase-04/business-rules.csv'),('Traceability','phase-04/traceability.csv'),('NFR Trace','phase-04/nonfunctional-trace.csv'),('Access','phase-04/access-matrix.csv'),('Stakeholders','phase-02/stakeholders.csv'),('Findings','phase-02/discovery-findings.csv'),('RACI','phase-02/responsibility-matrix.csv'),('AS-IS Steps','phase-03/as-is-step-register.csv'),('SAP Evidence','phase-05/sap-evidence.csv'),('Fit Gap','phase-06/fit-gap.csv'),('Backlog','phase-07/backlog.csv'),('UAT Cases','phase-09/uat-cases.csv'),('KPI Contract','phase-08/kpi-contract.csv'),('Data Model','phase-08/data-model.csv'),('Data Dictionary','phase-08/data-dictionary.csv'),('Risks','phase-10/risks.csv'),('Change Impact','phase-10/change-impact.csv'),('Business Case','phase-10/business-case-options.csv'),('Sensitivity','phase-10/sensitivity.csv')]
sheets += [(p.stem.title().replace('_',' '),'phase-08/data/'+p.name) for p in sorted((ROOT/'phase-08/data').glob('*.csv'))]
for title,path in sheets:
 rows=csvread(path);ws=wb.create_sheet(title[:31]);ws.append(list(rows[0]))
 for r in rows:
  ws.append([int(v) if re.fullmatch(r'-?\d+',v or '') else float(v) if re.fullmatch(r'-?\d+\.\d+',v or '') else v for v in r.values()])
 ws.auto_filter.ref=ws.dimensions
for ws in wb:
 ws.freeze_panes='A2';ws.sheet_view.showGridLines=False
 for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='112B42');cell.font=Font(color='FFFFFF',bold=True)
 for row in ws:
  for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
  ws.row_dimensions[row[0].row].height=95 if ws.title in ['UAT Cases','Requirements','Backlog','Fit Gap','KPI Contract','SAP Evidence'] else 45
 for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=20 if i==1 else 43
 ws.print_title_rows='1:1';ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
wb.save(OUT/'NorthStar_Complete_Workbook.xlsx')

# Word report: narrative, decisions and learning reviews; long catalogues live in the workbook.
report=Document();sec=report.sections[0];sec.header.paragraphs[0].text='NORTHSTAR | INDEPENDENT SIMULATED PORTFOLIO';sec.footer.paragraphs[0].text='AI-assisted draft · No SAP UAT executed · Learner review pending'
report.styles['Normal'].font.name='Calibri';report.styles['Normal'].font.size=Pt(10)
sec.top_margin=sec.bottom_margin=Inches(.65)
for n,p in enumerate(docpaths):
 if n:report.add_page_break()
 lines=p.read_text().splitlines();i=0
 while i<len(lines):
  line=lines[i]
  if line.startswith('```'):
   lang=line[3:];i+=1
   while i<len(lines) and not lines[i].startswith('```'):i+=1
   report.add_paragraph('Process flowchart: view the editable .mmd/.dot and rendered SVG in the relevant phase folder or offline viewer.' if lang=='mermaid' else 'See the editable Markdown source for code.')
  elif line.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].startswith('|'):
    if not sep(lines[i]):rows.append(cells(lines[i]))
    i+=1
   if len(rows[0])>5:report.add_paragraph('Wide reference table: see this document in the offline viewer or corresponding workbook/CSV for full readable detail.')
   else:
    t=report.add_table(rows=0,cols=len(rows[0]));t.style='Light Shading Accent 1'
    for r in rows:
     for c,v in zip(t.add_row().cells,r):c.text=v
   continue
  elif line.startswith('#'):report.add_heading(clean(line.lstrip('# ')),level=min(len(line)-len(line.lstrip('#')),3))
  elif line.strip():report.add_paragraph(clean(line.lstrip('> ')),style='List Bullet' if line.startswith('- ') else None)
  i+=1
report.save(OUT/'NorthStar_Portfolio_Report.docx')

slides=[
('NorthStar Manufacturing UK','SAP S/4HANA Order-to-Cash\nBusiness Process Analysis & Transformation',['Independent simulated portfolio','Public Edition 2602 research reference','Free documentation and synthetic-data route'], 'Explain that no real client, production configuration or implementation is represented.'),
('The business problem','Three connected operational problems',['Poor visibility → repeated enquiries and unclear ownership','Stock delays → unsupported promises and partial-delivery decisions','Manual billing handoffs → eligible deliveries waiting for review'], 'Consequences are hypotheses from the supplied fictional scenario, not real observed findings.'),
('Scope and representative order','One company · one warehouse · domestic B2B · GBP',['NS-ORD-1001: 100 brackets × £25 = £2,500 net','60 physically on hand; availability remains unresolved','Requested receipt 9 October; no confirmed promise invented','Order receipt through payment allocation and reporting'], 'Manufacturing planning, procurement redesign, international trade, complex tax, full migration and deployment are excluded.'),
('Discovery and process choices','Role-play reveals decisions, not stakeholder approval',['Validate data before confirmation; retain the original promise','Require agreement for partial delivery; keep remainder open','Finance owns credit release and billable-population review','Separate physical dispatch, goods-issue posting and receipt'], 'Explain alternative causes for late billing; email handoff is a hypothesis, not a proven root cause.'),
('Requirements to validation','One traceable evidence chain',['23 requirements: 3 business + 15 functional + 5 non-functional','15 stories and 12 rules with explicit exceptions','Example: PROB-002 → BR-002 → FR-004 → PC-004','US-004 → TC-005/006 → KPI-003'], 'All requirements and stories are proposed. Trace links exist but actual business validation is pending.'),
('SAP standard-first assessment','Documented capability is not demonstrated configuration',['Seven official sources reviewed; two pages explicitly unverified','BD9 stock-sales reference and payment handling researched','15 functional fit assessments remain desk-based','Validate exact roles, corrections, receipt capture and audit coverage'], 'Reference is Public Edition 2602; Learning pages are not release-pinned. Entitlements and configuration are not established.'),
('Synthetic baseline','Reproducible analysis, not achieved performance',[], 'Single-line orders, net-only finance, seeded data. On-time denominator includes overdue open orders. Cycle time excludes unfinished orders; billing mean excludes unbilled deliveries.'),
('Validation and remaining gates','Honest execution status',[f'{ctx["checks"]} synthetic data checks passed; six Python/SQL KPIs agree','22 SAP UAT cases designed; 0 executed','Desk walkthrough issues kept separate from SAP defects','Learner explanation, real discovery and tenant testing remain pending'], 'Do not turn data check success into a passed SAP UAT case or personal experience claim.'),
('Options and illustrative economics','Improve controls before assuming a large investment',['Compare manual controls, process standardisation and SAP-aligned design','6,000 orders × 5 minutes = 500 hours annual capacity','At £25/hour: £12,500 capacity equivalent, not cash savings','No cash ROI/payback claimed; costs are planning placeholders'], 'Recommend manual ownership/reconciliation first, then validate process standardisation. No SAP price quote or achieved benefit is represented.'),
('What this portfolio demonstrates','A reviewable foundation, with explicit limits',['Problem → evidence → requirements → design → validation reasoning','Reproducible synthetic analysis and source-labelled SAP research','Next: learner review, policy validation and system demonstration','AI drafted/generated artifacts; personal contribution must be stated accurately'], 'Final learning gate remains open until the learner can explain decisions without reading a script. No employment guarantee.'),
]
prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)
def textbox(slide,x,y,w,h,text,size=20,color='213547',bold=False):
 box=slide.shapes.add_textbox(PI(x),PI(y),PI(w),PI(h));tf=box.text_frame;tf.word_wrap=True
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.size=PP(size);p.font.color.rgb=RGBColor.from_string(color);p.font.bold=bold;p.space_after=PP(12)
 return box
for idx,(title,sub,bullets,notes) in enumerate(slides):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string('F3F6F8')
 textbox(sl,.55,.3,12,.3,'NORTHSTAR  /  INDEPENDENT SIMULATED PORTFOLIO',11,'087E8B',True)
 textbox(sl,.55,.9,12,1,title,32,'112B42',True);textbox(sl,.55,1.9,12,.9,sub,20,'526476')
 if idx==6:sl.shapes.add_picture(str(ROOT/'phase-08/kpi-baseline.png'),PI(.65),PI(2.8),width=PI(12))
 else:textbox(sl,.75,3,11.8,3.7,'\n'.join('• '+b for b in bullets),22)
 textbox(sl,.55,7,11,.25,'11 October 2026 · No production implementation or achieved savings claimed',10,'526476');textbox(sl,12,7,.7,.3,f'{idx+1:02}',11,'087E8B')
 sl.notes_slide.notes_text_frame.text=notes
prs.save(OUT/'NorthStar_Executive_Presentation.pptx')
(ROOT/'phase-11/executive-presentation.md').write_text('# Executive presentation — ten slides\n\nIndependent simulated portfolio; editable PPTX in downloads.\n\n'+'\n\n'.join(f'## {i+1}. {t}\n{s}\n\n'+'\n'.join('- '+b for b in bs)+'\n\nSpeaker note: '+n for i,(t,s,bs,n) in enumerate(slides)))
# Reopen editable formats; count substantive content, not just filenames.
assert len(Presentation(OUT/'NorthStar_Executive_Presentation.pptx').slides)==10
check=load_workbook(OUT/'NorthStar_Complete_Workbook.xlsx',read_only=True)
assert check['Requirements'].max_row==24 and check['UAT Cases'].max_row==23
assert len(Document(OUT/'NorthStar_Portfolio_Report.docx').paragraphs)>100
print(f'Built standalone viewer, {len(wb.sheetnames)}-sheet workbook, editable report and ten-slide presentation. Reopening checks passed.')
