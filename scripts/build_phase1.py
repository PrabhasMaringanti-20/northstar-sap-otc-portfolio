"""Build editable Phase 1 artifacts from Markdown sources.

Requires python-docx and openpyxl. No SAP access or external data is used.
"""
from pathlib import Path
import re
import zipfile
import csv
from html import escape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'downloads'
OUT.mkdir(exist_ok=True)

def clean(s):
    return re.sub(r'\[([^]]+)\]\([^)]+\)', r'\1', s).replace('**', '')

def cells(line):
    return [clean(c.strip()) for c in line.strip().strip('|').split('|')]

def separator(line):
    return all(re.fullmatch(r':?-+:?', c.strip()) for c in line.strip().strip('|').split('|'))

sources = [ROOT/'phase-01'/name for name in
           ['charter.md', 'delivery-plan.md', 'representative-order.md', 'learning-review.md']]
sources += [ROOT/'registers'/name for name in ['assumptions.md', 'decisions.md', 'glossary.md', 'progress.md']]
sources.append(ROOT/'phase-01/free-working-route.md')

doc = Document()
section = doc.sections[0]
section.header.paragraphs[0].text = 'NORTHSTAR MANUFACTURING UK | INDEPENDENT SIMULATED PORTFOLIO'
section.footer.paragraphs[0].text = 'Phase 1 • Draft for learner review • No production implementation or achieved benefits claimed'
section.top_margin = section.bottom_margin = Inches(.7)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10)
for source in sources:
    if source != sources[0]:
        doc.add_page_break()
    lines = source.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            block = []
            while i < len(lines) and lines[i].startswith('|'):
                if not separator(lines[i]): block.append(cells(lines[i]))
                i += 1
            table = doc.add_table(rows=0, cols=len(block[0]))
            table.style = 'Light Shading Accent 1'
            for row in block:
                for cell, value in zip(table.add_row().cells, row): cell.text = value
            continue
        if line.startswith('#'):
            level = len(line)-len(line.lstrip('#'))
            doc.add_heading(clean(line.lstrip('#').strip()), level=min(level, 3))
        elif line.startswith('- '): doc.add_paragraph(clean(line[2:]), style='List Bullet')
        elif line.strip(): doc.add_paragraph(clean(line))
        i += 1
docpath = OUT/'NorthStar_Phase_1_Charter_and_Learning_Pack.docx'
doc.save(docpath)

wb = Workbook()
guide = wb.active
guide.title = 'Read Me'
for row in [
    ['NorthStar Phase 1 working register', 'Draft v0.1 — 8 October 2026'],
    ['Evidence status', 'Independent simulation; no stakeholder approval, SAP demonstration or achieved savings'],
    ['Scope', 'Phase 1 only; all subsequent phases not started'],
    ['SAP reference', 'Public Edition provisionally selected; release and access unresolved'],
    ['Learning schedule', '4–6 hours/week assumed, not confirmed'],
    ['Editing', 'Edit Markdown source files to rebuild consistently; reconcile workbook edits before regeneration'],
    ['Verification', 'Artifact structure checks are not SAP UAT or validation of business outcomes'],
    ['Free working route', 'Use LibreOffice or Markdown/CSV; no paid SAP access, Office licence or hosting required'],
]: guide.append(row)

sheet_specs = [
    ('Problems','phase-01/charter.md',0),
    ('Objectives','phase-01/charter.md',1),
    ('Stakeholders','phase-01/charter.md',2),
    ('Delivery Plan','phase-01/delivery-plan.md',0),
    ('Representative Order','phase-01/representative-order.md',0),
    ('Assumptions','registers/assumptions.md',0),
    ('Decisions','registers/decisions.md',0),
    ('Glossary','registers/glossary.md',0),
    ('Progress','registers/progress.md',0),
    ('Open Questions','registers/progress.md',1),
]
for title, relative, index in sheet_specs:
    blocks = []
    active = []
    for line in (ROOT/relative).read_text().splitlines()+['']:
        if line.startswith('|'):
            if not separator(line): active.append(cells(line))
        elif active:
            blocks.append(active)
            active = []
    ws = wb.create_sheet(title)
    for row in blocks[index]: ws.append(row)
    ws.auto_filter.ref = ws.dimensions

for ws in wb:
    ws.freeze_panes = 'A2'
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[row[0].row].height = 65 if ws.title == 'Read Me' else 76
    for col in range(1, ws.max_column+1):
        ws.column_dimensions[get_column_letter(col)].width = 26 if col == 1 else 48
    ws.sheet_view.showGridLines = False
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
bookpath = OUT/'NorthStar_Phase_1_Working_Registers.xlsx'
wb.save(bookpath)

csv_dir = OUT/'csv'
csv_dir.mkdir(exist_ok=True)
for ws in wb:
    with (csv_dir/(ws.title.replace(' ', '_')+'.csv')).open('w', newline='', encoding='utf-8-sig') as f:
        csv.writer(f).writerows(ws.values)

# Standalone reading/printing artifact; no network dependencies or executable scripts.
parts = []
nav = []
for number, source in enumerate(sources):
    anchor = f'section-{number}'
    title = source.read_text().splitlines()[0].lstrip('# ')
    nav.append(f'<a href="#{anchor}">{escape(title)}</a>')
    parts.append(f'<section id="{anchor}">')
    lines = source.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not separator(lines[i]): rows.append(cells(lines[i]))
                i += 1
            parts.append('<div class="table-wrap"><table><thead>')
            for idx, row in enumerate(rows):
                tag = 'th' if idx == 0 else 'td'
                parts.append('<tr>'+''.join(f'<{tag}>{escape(v)}</{tag}>' for v in row)+'</tr>')
                if idx == 0: parts.append('</thead><tbody>')
            parts.append('</tbody></table></div>')
            continue
        if line.startswith('#'):
            level = min(3, len(line)-len(line.lstrip('#'))+1)
            parts.append(f'<h{level}>{escape(clean(line.lstrip("# ")))}</h{level}>')
        elif line.strip():
            parts.append('<p>'+escape(clean(line))+'</p>')
        i += 1
    parts.append('</section>')
css = '''body{margin:0;background:#f2f5f8;color:#203047;font:16px/1.6 system-ui,sans-serif}
header,main{max-width:1100px;margin:auto;padding:28px}header{background:#17365d;color:white;box-sizing:border-box}
h1{line-height:1.2}h2,h3{color:#17365d}nav{display:flex;flex-wrap:wrap;gap:10px;margin:20px 0}
nav a{padding:8px 12px;background:white;border:1px solid #cad5e2;border-radius:5px;color:#17365d}
section{background:white;padding:28px;margin:22px 0;border:1px solid #dce3eb;border-radius:8px}
.table-wrap{overflow:auto}table{border-collapse:collapse;width:100%;font-size:14px}th,td{padding:10px;border:1px solid #cbd5e1;text-align:left;vertical-align:top}
th{background:#e9eff6}tr:nth-child(even){background:#f7f9fc}a:focus{outline:3px solid #ac6200}
@media print{body{background:white;font-size:10pt}header,main{padding:0;max-width:none}header{background:white;color:black}nav{display:none}section{border:0;padding:0;break-before:page}.table-wrap{overflow:visible}tr{break-inside:avoid}th{background:#eee}}
@media(max-width:600px){main,header{padding:16px}section{padding:14px}}
'''
html = '<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>NorthStar | Phase 1 portfolio pack</title><style>'+css+'</style></head><body><header><p>INDEPENDENT SIMULATED PORTFOLIO · PHASE 1</p><h1>NorthStar Manufacturing UK</h1><p>Order-to-Cash analysis · Draft for learner review · No achieved business benefits claimed</p></header><main><p>Free working route: read offline, edit the supplied sources, or print this document to PDF using your browser.</p><nav aria-label="Document sections">'+''.join(nav)+'</nav>'+''.join(parts)+'</main></body></html>'
(OUT/'NorthStar_Phase_1.html').write_text(html, encoding='utf-8')

# Confirm generated artifacts can be reopened and expected content survived.
check = load_workbook(bookpath)
assert len(check.sheetnames) == 11
assert check['Problems']['A2'].value == 'PROB-001'
assert check['Representative Order']['B2'].value.startswith('NS-ORD-1001')
assert len(Document(docpath).tables) == 10

archive = OUT/'NorthStar_Phase_1_Portfolio_Pack.zip'
with zipfile.ZipFile(archive, 'w', zipfile.ZIP_DEFLATED) as z:
    for p in sorted(ROOT.rglob('*')):
        if p.is_file() and p.suffix != '.zip' and '__pycache__' not in p.parts:
            z.write(p, Path('northstar-otc')/p.relative_to(ROOT))
print('Built and reopened Word and Excel artifacts; 11 workbook sheets and 10 document tables checked.')
print('Created ZIP with editable sources and reproducible build script.')
