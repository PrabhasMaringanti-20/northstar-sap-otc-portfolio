"""Complete free-tool local build. No publication, network calls or SAP interaction."""
from pathlib import Path
import subprocess,sys,json,zipfile,hashlib
ROOT=Path(__file__).resolve().parents[1]
for script in ['build_data.py','update_project_state.py','build_portfolio.py','verify_project.py','build_outputs.py']:
 subprocess.run([sys.executable,str(ROOT/'scripts'/script)],check=True,cwd=ROOT)
subprocess.run(['node','--check',str(ROOT/'viewer/app.js')],check=True)
subprocess.run(['node',str(ROOT/'scripts/check_viewer.js')],check=True)

from openpyxl import load_workbook
from docx import Document
from pptx import Presentation
from html.parser import HTMLParser
out=ROOT/'downloads';html=(ROOT/'viewer/index.html').read_text()
class DependencyCheck(HTMLParser):
 def __init__(self):super().__init__();self.external=[]
 def handle_starttag(self,tag,attrs):
  d=dict(attrs)
  if tag in ['script','img','iframe','link'] and (d.get('src') or d.get('href')):self.external.append((tag,d))
p=DependencyCheck();p.feed(html);assert not p.external,'Viewer must not load external dependencies'
assert len(Presentation(out/'NorthStar_Executive_Presentation.pptx').slides)==10
assert len(load_workbook(out/'NorthStar_Complete_Workbook.xlsx',read_only=True).sheetnames)==31
assert len(Document(out/'NorthStar_Portfolio_Report.docx').paragraphs)>100
from browser_evidence import current_browser_evidence
browser=current_browser_evidence()
artifact=dict(date='2026-10-11',office_reopening='Passed: DOCX, 31-sheet XLSX, 10-slide PPTX',offline_dependencies='Passed: no external rendering dependencies',viewer_logic=json.loads((ROOT/'quality/viewer-checks.json').read_text()),browser_evidence='quality/browser-checks.json' if browser else 'Pending current-source browser test',visual_limit='See real-browser evidence and exact tested HTML hash; LibreOffice visual rendering unverified' if browser else 'Current-source browser and LibreOffice rendering unverified',sap_uat='22 Not executed')
(ROOT/'quality/artifact-checks.json').write_text(json.dumps(artifact,indent=2))
archive=out/'NorthStar_Complete_Portfolio.zip'
with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
 for f in sorted(ROOT.rglob('*')):
  if f.is_file() and f.suffix not in ['.zip','.sha256'] and '__pycache__' not in f.parts and '.git' not in f.parts:
   # Historical Phase 1 downloads remain accessible separately; avoid duplicating stale snapshots in the full pack.
   if f.parent==out and f.name.startswith('NorthStar_Phase_1'):continue
   if 'downloads/csv/' in str(f):continue
   z.write(f,Path('northstar-otc')/f.relative_to(ROOT))
with zipfile.ZipFile(archive) as z:assert z.testzip() is None
(out/'NorthStar_Complete_Portfolio.sha256').write_text(hashlib.sha256(archive.read_bytes()).hexdigest()+'  '+archive.name+'\n')
subprocess.run([sys.executable,str(ROOT/'scripts/build_handover.py')],check=True,cwd=ROOT)
subprocess.run([sys.executable,str(ROOT/'scripts/package_release.py')],check=True,cwd=ROOT)
print('Complete portfolio and PDF handover built; no publishing or deployment performed.')
